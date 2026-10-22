from openpyxl import load_workbook

from bdl_report.builder import build
from bdl_report.cli import main
from bdl_report.client import BdlClient
from bdl_report.demo import GMINA, DemoSession
from bdl_report.indicators import DEFAULT_INDICATORS, Indicator, load_indicators
from bdl_report.render import fmt

YEARS = list(range(2014, 2025))


def demo_report():
    client = BdlClient(session=DemoSession(), min_interval=0)
    return build(client, GMINA, DEFAULT_INDICATORS, YEARS)


def test_parents_are_resolved():
    report = demo_report()
    assert report.powiat.name.startswith("Powiat")
    assert report.voivodeship.name.startswith("Województwo")


def test_missing_gmina_data_falls_back_to_powiat():
    by_id = {s.indicator.id: s for s in demo_report().series}
    assert by_id[72305].source is None  # gmina has its own population data
    assert by_id[64428].source.level == 5  # wages only exist for the powiat
    assert by_id[64428].years == YEARS


def test_comparisons_only_for_comparable_indicators():
    by_id = {s.indicator.id: s for s in demo_report().series}
    assert set(by_id[60559].comparisons) == {"Województwo przykładowe (DEMO)", "Polska"}
    assert by_id[72305].comparisons == {}


def test_change_pct():
    population = demo_report().series[0]
    first, last = population.values[2014], population.values[2024]
    assert population.change_pct == (last / first - 1) * 100


def test_unknown_indicator_has_no_data_but_report_still_builds():
    client = BdlClient(session=DemoSession(), min_interval=0)
    report = build(client, GMINA, [Indicator(1, "Nieistniejąca")], YEARS)
    assert report.series[0].values == {}
    assert report.series[0].variable is None


def test_fmt_uses_polish_separators():
    assert fmt(6155.23) == "6 155"
    assert fmt(5.6) == "5,6"
    assert fmt(None) == "–"


def test_load_custom_indicators(tmp_path):
    path = tmp_path / "z.json"
    path.write_text('[{"id": 123, "label": "Mieszkania", "compare": true}]', encoding="utf-8")
    assert load_indicators(path) == [Indicator(123, "Mieszkania", compare=True)]


def test_cli_demo_writes_html_and_excel(tmp_path, capsys):
    assert (
        main(
            ["--demo", "raport", "Przykladowo", "-o", str(tmp_path), "--od", "2014", "--do", "2024"]
        )
        == 0
    )
    html = next(tmp_path.glob("*.html")).read_text(encoding="utf-8")
    assert "Przykładowo (DEMO)" in html and "data:image/png;base64," in html
    wb = load_workbook(next(tmp_path.glob("*.xlsx")))
    assert wb.sheetnames[0] == "Podsumowanie"
    assert len(wb.sheetnames) == 5
    assert "Raport:" in capsys.readouterr().out


def test_cli_search_and_variables(capsys):
    assert main(["--demo", "szukaj", "przykład"]) == 0
    assert "099999999996" in capsys.readouterr().out
    assert main(["--demo", "zmienne", "--sprawdz"]) == 0
    assert "GUS: Ludność (DEMO): ogółem [osoba] (najniżej: gmina)" in capsys.readouterr().out


def test_report_file_names_are_ascii(tmp_path):
    main(["--demo", "raport", "x", "-o", str(tmp_path)])
    assert {p.name for p in tmp_path.iterdir()} == {
        "przykladowo_demo.html",
        "przykladowo_demo.xlsx",
    }


def test_excel_layout(tmp_path):
    main(["--demo", "raport", "x", "-o", str(tmp_path), "--od", "2014", "--do", "2024"])
    wb = load_workbook(tmp_path / "przykladowo_demo.xlsx")
    ws = wb["Podsumowanie"]
    assert ws["B1"].value == "Przykładowo (DEMO)"
    assert ws.sheet_view.showGridLines is False
    assert wb.properties.creator == "bdl-report"
    sheet = wb["Przeciętne wynagrodzenie brutto"]
    assert sheet["B5"].value == "Rok" and sheet["C5"].value == "Powiat przykładowski (DEMO)"
    assert sheet["C6"].number_format == "#,##0"
    assert len(sheet._charts) == 1


def test_variable_search_goes_through_subjects(capsys):
    assert main(["--demo", "zmienne", "wynagrodzenia"]) == 0
    out = capsys.readouterr().out
    assert "64428" in out and "Przeciętne miesięczne wynagrodzenia brutto" in out


def test_variable_search_level_filter(capsys):
    assert main(["--demo", "zmienne", "DEMO", "--poziom", "gmina"]) == 0
    out = capsys.readouterr().out
    assert "72305" in out and "64428" not in out  # wages only go down to powiat


def test_methodology_break_is_flagged():
    from bdl_report.builder import Series

    s = Series(Indicator(1, "Gęstość"), None, {2021: 1894.8, 2022: 1859.2, 2023: 1149.2})
    assert [y for y, _ in s.breaks] == [2023]
    assert "2023 (-38%)" in s.break_note
    assert Series(Indicator(1, "x"), None, {2020: 100, 2021: 110}).break_note == ""
