"""Excel workbook: a summary sheet first, then one sheet per indicator with data and a chart.

Same visual language as the HTML page: navy and gold, quiet greys, real number formats,
ready to print on A4.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .builder import ReportData, Series
from .render import series_title

NAVY = "1F4E79"
GOLD = "C9A227"
GREY = "8C8C8C"
INK = "1B2430"
MUTED = "6B7280"
LINE = "D9DEE5"
BAND = "F5F7FA"
GOOD = "1E6B3A"
BAD = "9B1C1C"
FONT = "Calibri"
SERIES_COLORS = [NAVY, GOLD, GREY]  # unit, voivodeship, Poland

_ROW_BORDER = Border(bottom=Side(style="thin", color=LINE))
_HEADER_BORDER = Border(bottom=Side(style="medium", color=GOLD))
_CHANGE_FORMAT = "+0.0%;-0.0%;0.0%"


def _font(size: float = 10, bold: bool = False, color: str = INK, **kw) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, **kw)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def number_format(s: Series) -> str:
    """Thousands separator for big numbers, one decimal for rates such as 5,6 %."""
    values = list(s.values.values()) + [v for c in s.comparisons.values() for v in c.values()]
    if not values or max(abs(v) for v in values) >= 1000:
        return "#,##0"
    if all(float(v).is_integer() for v in values):
        return "#,##0"
    return "#,##0.0"


def _unit_label(s: Series) -> str:
    return s.variable.unit if s.variable and s.variable.unit else ""


def _title(ws: Worksheet, title: str, subtitle: str, last_col: int) -> None:
    ws.cell(1, 2, title).font = Font(name="Calibri Light", size=18, color=NAVY)
    ws.cell(2, 2, subtitle).font = _font(10, color=MUTED)
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[2].height = 18
    for col in range(2, last_col + 1):
        ws.cell(3, col).border = Border(bottom=Side(style="medium", color=GOLD))
    ws.row_dimensions[3].height = 6
    ws.column_dimensions["A"].width = 2


def _header(ws: Worksheet, row: int, headers: list[str], first_col: int = 2, right=()) -> None:
    for i, text in enumerate(headers):
        cell = ws.cell(row, first_col + i, text)
        cell.font = _font(10, bold=True, color="FFFFFF")
        cell.fill = _fill(NAVY)
        cell.border = _HEADER_BORDER
        cell.alignment = Alignment(
            horizontal="right" if i in right else None, vertical="center", wrap_text=True
        )
    ws.row_dimensions[row].height = 30


def _section(ws: Worksheet, row: int, text: str) -> None:
    ws.cell(row, 2, text).font = _font(11, bold=True, color=NAVY)
    ws.row_dimensions[row].height = 22


def _page_setup(ws: Worksheet, footer: str, landscape: bool = False) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.print_options.horizontalCentered = True
    ws.oddFooter.left.text = footer
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Strona &P z &N"
    ws.oddFooter.right.size = 8


def _summary(ws: Worksheet, report: ReportData, footer: str) -> None:
    unit = report.unit
    where = ", ".join(u.name for u in report.ancestors)
    subtitle = " · ".join(
        p
        for p in (
            unit.level_name,
            where,
            f"lata {report.years[0]}–{report.years[-1]}",
            "Bank Danych Lokalnych GUS",
        )
        if p
    )
    widths = {"B": 40, "C": 16, "D": 13, "E": 8, "F": 11, "G": 13, "H": 36}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    _title(ws, unit.name, subtitle, last_col=8)

    row = 5
    _section(ws, row, "Wskaźniki")
    row += 1
    _header(
        ws,
        row,
        ["Wskaźnik", "Ostatnia wartość", "Jednostka", "Rok", "Zmiana", "Okres", "Uwagi"],
        right={1, 3, 4},
    )
    ws.cell(row, 7).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row, 4).alignment = Alignment(vertical="center", indent=2)
    for n, s in enumerate(report.series, start=1):
        row += 1
        latest, change = s.latest, s.change_pct
        note = f"dane: {s.source.name}" if s.source else ""
        if not s.values:
            note = "brak danych w BDL dla tej jednostki i lat"
        values = [
            series_title(s),
            latest[1] if latest else None,
            _unit_label(s) or None,
            latest[0] if latest else None,
            change / 100 if change is not None else None,
            f"{s.years[0]}–{s.years[-1]}" if len(s.years) > 1 else None,
            note or None,
        ]
        for i, value in enumerate(values):
            cell = ws.cell(row, 2 + i, value)
            cell.font = _font(10)
            cell.border = _ROW_BORDER
            cell.alignment = Alignment(
                horizontal={1: "right", 3: "right", 4: "right", 5: "center"}.get(i),
                vertical="center",
                wrap_text=i in (0, 6),
                indent=2 if i == 2 else 0,
            )
            if n % 2 == 0:
                cell.fill = _fill(BAND)
        ws.cell(row, 3).number_format = number_format(s)
        ws.cell(row, 3).font = _font(10, bold=True)
        ws.cell(row, 5).number_format = "0"
        change_cell = ws.cell(row, 6)
        change_cell.number_format = _CHANGE_FORMAT
        better = s.indicator.higher_is_better
        if change is not None and better is not None and abs(change) >= 0.05:
            change_cell.font = _font(10, bold=True, color=GOOD if (change > 0) == better else BAD)
        ws.cell(row, 8).font = _font(9, color=MUTED, italic=True)
        ws.row_dimensions[row].height = 22

    compared = [s for s in report.series if s.comparisons and s.values]
    if compared:
        row += 2
        _section(ws, row, "Na tle regionu i kraju")
        row += 1
        # One column per indicator (latest year shared by all three units), one row per unit.
        columns = []
        for s in compared:
            year = max(set(s.values).intersection(*s.comparisons.values()), default=None)
            if year is not None:
                columns.append((s, year))
        _header(
            ws,
            row,
            ["", *(f"{series_title(s)} ({year})" for s, year in columns)],
            right=set(range(1, len(columns) + 1)),
        )
        ws.row_dimensions[row].height = 44
        own = compared[0].source.name if compared[0].source else unit.name
        units = [own, *compared[0].comparisons]
        for n, name in enumerate(units, start=1):
            row += 1
            label = ws.cell(row, 2, name)
            label.font = _font(10, bold=(n == 1))
            label.border = _ROW_BORDER
            label.alignment = Alignment(vertical="center")
            for i, (s, year) in enumerate(columns, start=3):
                value = s.values[year] if n == 1 else s.comparisons.get(name, {}).get(year)
                cell = ws.cell(row, i, value)
                cell.font = _font(10, bold=(n == 1))
                cell.number_format = number_format(s)
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.border = _ROW_BORDER
            if n % 2 == 0:
                for col in range(2, 3 + len(columns)):
                    ws.cell(row, col).fill = _fill(BAND)
            ws.row_dimensions[row].height = 22

    row += 2
    _section(ws, row, "Źródła i definicje")
    for s in report.series:
        if not s.variable:
            continue
        row += 1
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        unit_txt = f" [{s.variable.unit}]" if s.variable.unit else ""
        cell = ws.cell(
            row,
            2,
            f"{series_title(s)}: zmienna BDL {s.indicator.id}, „{s.variable.name}”{unit_txt}",
        )
        cell.font = _font(9, color=MUTED)
    row += 2
    ws.merge_cells(start_row=row, start_column=2, end_row=row + 1, end_column=8)
    note = ws.cell(
        row,
        2,
        f"Źródło: Bank Danych Lokalnych GUS (bdl.stat.gov.pl), jednostka {unit.id}. "
        f"Wygenerowano {datetime.now():%d.%m.%Y %H:%M}. GUS publikuje większość danych "
        "z opóźnieniem około roku, dlatego ostatni rok bywa pusty.",
    )
    note.font = _font(9, color=MUTED, italic=True)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 18
    ws.row_dimensions[row + 1].height = 18
    _page_setup(ws, footer)


def _series_sheet(ws: Worksheet, report: ReportData, s: Series, footer: str) -> None:
    names = [s.source.name if s.source else report.unit.name, *s.comparisons]
    all_series = [s.values, *s.comparisons.values()]
    years = sorted(set().union(*all_series))
    unit_txt = _unit_label(s)
    subtitle = " · ".join(
        p
        for p in (
            f"GUS: {s.variable.name}" if s.variable else "",
            f"jednostka miary: {unit_txt}" if unit_txt else "",
            f"dane powiatu ({s.source.name})" if s.source else "",
        )
        if p
    )
    last_col = 2 + len(names)
    ws.column_dimensions["B"].width = 8
    for i in range(len(names)):
        ws.column_dimensions[get_column_letter(3 + i)].width = 24
    _title(ws, series_title(s), subtitle, last_col=max(last_col, 10))

    header_row = 5
    _header(ws, header_row, ["Rok", *names], right=set(range(len(names) + 1)))
    fmt = number_format(s)
    for n, year in enumerate(years, start=1):
        row = header_row + n
        values = [year, *(series.get(year) for series in all_series)]
        for i, value in enumerate(values):
            cell = ws.cell(row, 2 + i, value)
            cell.font = _font(10, bold=(i == 1))
            cell.border = _ROW_BORDER
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "0" if i == 0 else fmt
            if n % 2 == 0:
                cell.fill = _fill(BAND)
        ws.row_dimensions[row].height = 18
    last_row = header_row + len(years)
    ws.freeze_panes = ws.cell(header_row + 1, 2)

    chart = LineChart()
    chart.title = None
    chart.style = 2
    chart.height, chart.width = 8.5, 17
    if len(names) > 1:
        chart.legend.position = "b"
    else:
        chart.legend = None
    chart.y_axis.numFmt = fmt
    chart.y_axis.majorGridlines.spPr = None
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    data = Reference(ws, min_col=3, max_col=last_col, min_row=header_row, max_row=last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=header_row + 1, max_row=last_row))
    for line, color in zip(chart.series, SERIES_COLORS, strict=False):
        line.graphicalProperties.line.solidFill = color
        line.graphicalProperties.line.width = 28575  # 2.25 pt
        line.marker.symbol = "circle"
        line.marker.size = 5
        line.marker.graphicalProperties.solidFill = color
        line.marker.graphicalProperties.line.solidFill = color
        line.smooth = False
    ws.add_chart(chart, f"B{last_row + 3}")
    _page_setup(ws, footer)


def write_xlsx(report: ReportData, path: str | Path) -> None:
    footer = f"{report.unit.name} · Bank Danych Lokalnych GUS"
    wb = Workbook()
    wb.properties.creator = "bdl-report"
    wb.properties.title = f"{report.unit.name}: raport GUS"
    summary = wb.active
    summary.title = "Podsumowanie"
    _summary(summary, report, footer)
    for s in report.series:
        if s.values:
            _series_sheet(wb.create_sheet(_sheet_name(series_title(s))), report, s, footer)
    wb.active = 0
    wb.save(path)


def _sheet_name(title: str) -> str:
    for ch in "[]:*?/\\":
        title = title.replace(ch, " ")
    return title[:31]
